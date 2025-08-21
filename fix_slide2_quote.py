#!/usr/bin/env python3
"""
Script to specifically fix the quote on slide 2
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def fix_slide2_quote():
    """Fix the quote visibility on slide 2"""
    ppt_path = "/Users/alexanderfedin/Projects/nolock.social/marketing/REBUILD TRUST IN THE DIGITAL SPACEN LOCK•SOCIAL.pptx"
    
    # Load the presentation
    print(f"Loading presentation: {ppt_path}")
    prs = Presentation(ppt_path)
    
    # Get slide 2
    slide = prs.slides[1]  # 0-indexed
    
    # Quote content
    quote = "76% of users struggle to identify authentic content in digital spaces"
    
    # Identify all shapes on the slide
    shapes = list(slide.shapes)
    print(f"Found {len(shapes)} shapes on slide 2")
    
    # Find and fix the quote textbox (shape 15)
    quote_shape = None
    for shape in shapes:
        if hasattr(shape, 'text_frame') and quote in shape.text_frame.text:
            quote_shape = shape
            print(f"Found quote in shape: {shape.name}")
            # Fix position and size
            shape.left = Inches(1)
            shape.top = Inches(6)
            shape.width = Inches(8)
            shape.height = Inches(0.75)
            print("Restored quote visibility")
            break
    
    # If quote shape not found, create a new one
    if not quote_shape:
        print("Quote not found, creating new quote textbox")
        quote_textbox = slide.shapes.add_textbox(
            Inches(1),      # left
            Inches(6),      # top
            Inches(8),      # width
            Inches(0.75)    # height
        )
        quote_textbox.text_frame.text = f"\"{quote}\""
        # Format paragraph
        p = quote_textbox.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.runs[0]
        run.font.italic = True
        run.font.size = Pt(11)
    
    # Save the updated presentation
    print("Saving updated presentation")
    prs.save(ppt_path)
    print("Slide 2 quote fixed!")
    
    return True

def open_powerpoint():
    """Open PowerPoint and the presentation"""
    ppt_path = "/Users/alexanderfedin/Projects/nolock.social/marketing/REBUILD TRUST IN THE DIGITAL SPACEN LOCK•SOCIAL.pptx"
    os.system(f"open -a 'Microsoft PowerPoint' '{ppt_path}'")
    print("PowerPoint opened with the presentation")

if __name__ == "__main__":
    import argparse
    
    parser = argparse.ArgumentParser(description='Fix quote on slide 2')
    parser.add_argument('--open', action='store_true', help='Open PowerPoint after update')
    
    args = parser.parse_args()
    
    # Fix the quote
    success = fix_slide2_quote()
    
    # Open PowerPoint if requested
    if args.open:
        open_powerpoint()
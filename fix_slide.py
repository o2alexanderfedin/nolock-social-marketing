#!/usr/bin/env python3
"""
Script to fix a single PowerPoint slide by cleaning and updating content
"""

import os
import re
import sys
import argparse
from pptx import Presentation
import datetime
import subprocess

def extract_markdown_content(markdown_path):
    """Extract title, subtitle, and content from markdown file"""
    with open(markdown_path, 'r') as f:
        content = f.read()
    
    # Extract title (first # heading)
    title_match = re.search(r'^# (.+?)$', content, re.MULTILINE)
    title = title_match.group(1) if title_match else None
    
    # Extract subtitle (first ## heading)
    subtitle_match = re.search(r'^## (.+?)$', content, re.MULTILINE)
    subtitle = subtitle_match.group(1) if subtitle_match else None
    
    # Extract summary (content after subtitle)
    summary = ""
    bullet_points = []
    
    # Try to extract summary if it exists in export files
    export_path = f"/Users/alexanderfedin/Projects/nolock.social/marketing/slides_export/slide{int(markdown_path.split('slide')[-1].split('.')[0]):02d}_export.md"
    if os.path.exists(export_path):
        with open(export_path, 'r') as f:
            export_content = f.read()
            
        summary_match = re.search(r'## Summary\n(.*?)(?=\n\n|\Z)', export_content, re.DOTALL)
        if summary_match:
            summary = summary_match.group(1).strip()
            
        bullet_match = re.search(r'## Bullet Points\n(.*?)(?=\n\n|\Z)', export_content, re.DOTALL)
        if bullet_match:
            bullet_text = bullet_match.group(1).strip()
            bullet_points = [line.strip() for line in bullet_text.split('\n') if line.strip()]
    
    # If no summary from export, extract from markdown content (after subtitle)
    if not summary and subtitle_match:
        lines = content.split('\n')
        subtitle_index = next((i for i, line in enumerate(lines) if line.startswith('## ') and subtitle in line), -1)
        if subtitle_index >= 0 and subtitle_index + 1 < len(lines):
            # Extract text after subtitle until the next section or end
            remaining_lines = []
            for line in lines[subtitle_index + 1:]:
                if line.startswith('#') or line.startswith('---'):
                    break
                if line.strip() and not line.startswith('['):  # Skip navigation links
                    remaining_lines.append(line)
            
            summary = '\n'.join(remaining_lines).strip()
            
            # Extract bullet points
            bullet_pattern = r'^\s*[-*•]\s+(.+)$'
            bullet_points = [re.match(bullet_pattern, line).group(1) for line in remaining_lines 
                             if re.match(bullet_pattern, line)]
    
    return {
        'title': title,
        'subtitle': subtitle,
        'summary': summary,
        'bullet_points': bullet_points
    }

def clean_and_update_slide(slide_num):
    """Clean and update a specific slide in the PowerPoint presentation"""
    # File paths
    ppt_path = "/Users/alexanderfedin/Projects/nolock.social/marketing/REBUILD TRUST IN THE DIGITAL SPACEN LOCK•SOCIAL.pptx"
    formatted_num = f"{slide_num:02d}"
    md_path = f"/Users/alexanderfedin/Projects/nolock.social/marketing/pitch-deck-investor-full/slides/slide{formatted_num}.md"
    
    # Create backup
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = f"/Users/alexanderfedin/Projects/nolock.social/marketing/slide_{formatted_num}_backup_{timestamp}.pptx"
    
    # Check if markdown file exists
    if not os.path.exists(md_path):
        print(f"Error: Markdown file not found: {md_path}")
        return False
    
    # Extract content from markdown
    content = extract_markdown_content(md_path)
    print(f"Extracted content for slide {slide_num}:")
    print(f"  Title: {content['title']}")
    print(f"  Subtitle: {content['subtitle']}")
    print(f"  Summary length: {len(content['summary'])} chars")
    print(f"  Bullet points: {len(content['bullet_points'])}")
    
    # Load the presentation
    print(f"Loading presentation: {ppt_path}")
    prs = Presentation(ppt_path)
    
    # Save backup
    print(f"Saving backup to: {backup_path}")
    prs.save(backup_path)
    
    # Get the slide (0-indexed)
    if slide_num > len(prs.slides):
        print(f"Error: Slide {slide_num} does not exist (total slides: {len(prs.slides)})")
        return False
    
    slide = prs.slides[slide_num - 1]
    
    # Clean and update the slide
    print(f"Updating slide {slide_num}...")
    
    # Find shapes for title, subtitle, and content
    shapes = list(slide.shapes)
    title_shape = None
    subtitle_shape = None
    content_shapes = []
    
    for shape in shapes:
        if hasattr(shape, 'text'):
            # Simple heuristic: first shape with text is title, second is subtitle
            if title_shape is None:
                title_shape = shape
            elif subtitle_shape is None:
                subtitle_shape = shape
            else:
                content_shapes.append(shape)
    
    # Update title and subtitle
    if title_shape and content['title']:
        print(f"  - Setting title: '{content['title']}'")
        title_shape.text = content['title']
    
    if subtitle_shape and content['subtitle']:
        print(f"  - Setting subtitle: '{content['subtitle']}'")
        subtitle_shape.text = content['subtitle']
    
    # Update content (summary or bullet points)
    if content_shapes and (content['summary'] or content['bullet_points']):
        main_content_shape = content_shapes[0]
        
        if content['bullet_points']:
            # Use bullet points if available
            bullet_text = '\n'.join([f"• {point}" for point in content['bullet_points']])
            print(f"  - Setting bullet points ({len(content['bullet_points'])} items)")
            main_content_shape.text = bullet_text
        elif content['summary']:
            # Use summary if available
            print(f"  - Setting summary text ({len(content['summary'])} chars)")
            main_content_shape.text = content['summary']
    
    # Save the updated presentation
    print(f"Saving updated presentation")
    prs.save(ppt_path)
    print(f"Slide {slide_num} updated successfully!")
    
    return True

def open_powerpoint():
    """Open PowerPoint and the presentation"""
    ppt_path = "/Users/alexanderfedin/Projects/nolock.social/marketing/REBUILD TRUST IN THE DIGITAL SPACEN LOCK•SOCIAL.pptx"
    os.system(f"open -a 'Microsoft PowerPoint' '{ppt_path}'")
    print(f"PowerPoint opened with the presentation")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description='Fix a specific PowerPoint slide')
    parser.add_argument('slide_num', type=int, help='Slide number to fix')
    parser.add_argument('--open', action='store_true', help='Open PowerPoint after update')
    
    args = parser.parse_args()
    
    # Update the slide
    success = clean_and_update_slide(args.slide_num)
    
    # Open PowerPoint if requested and update was successful
    if success and args.open:
        open_powerpoint()